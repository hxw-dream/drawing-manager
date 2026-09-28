"""清册归类与 Excel 导出测试。"""
from io import BytesIO

import openpyxl

from app.inventory import CATEGORIES, categorize


def test_categorize_priority():
    assert categorize("设计说明 DESIGN SPECIFICATION") == "设计说明"
    assert categorize("THE LIST OF MAIN EQUIPMENT AND MATERIAL 材料设备清册") == "材料设备清册"
    assert categorize("行政办公楼屋顶避雷带布置图") == "防雷接地"
    assert categorize("CHEMICAL STORAGE ROOM CABLE TRAY LAYOUT") == "电缆敷设"
    assert categorize("取水泵房照明布置图") == "照明安装"
    assert categorize("取水PC配电室低压配电装置和变压器安装图") == "配电装置安装"
    assert categorize("某某接线图 WIRING") == "电气接线"
    assert categorize("不知什么图") == "其他"
    # 优先级：说明/清册 在布局类词之前
    assert categorize("电缆清册布置图") == "材料设备清册"


def test_categories_cover_expected_disciplines():
    names = [c for c, _ in CATEGORIES]
    assert {"设计说明", "材料设备清册", "防雷接地", "电缆敷设", "照明安装"} <= set(names)


def test_export_xlsx_sheets(tmp_db, sample_library):
    from app import inventory as inv_mod
    from app import scanner
    from app.db import connect
    with connect() as conn:
        scanner.scan_root(str(sample_library), log=lambda m: None)
        inv = inv_mod.build_inventory(conn, project=1)
        data = inv_mod.make_xlsx(inv)
    wb = openpyxl.load_workbook(BytesIO(data))
    names = wb.sheetnames
    assert "tuzhi" in names            # 每个导入文件夹一张清册表
    assert "图纸明细" in names and "分类统计" in names
    ws = wb["tuzhi"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "序号"        # ZX 样式表头
    assert any(r[1] == "D0401" for r in rows[1:])
